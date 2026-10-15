import os
import sys
import re
import uuid
import threading
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from ebooklib import epub
from bs4 import BeautifulSoup
import customtkinter as ctk
from tkinter import filedialog, messagebox, simpledialog

# Set up CustomTkinter appearance
ctk.set_appearance_mode("System")
ctk.set_default_color_theme("blue")

spinner = ["/", "-", "\\", "|"]

LANGUAGE_OPTIONS = [
    "English (US) [en-US]",
    "English (UK) [en-GB]",
    "English (Australia) [en-AU]",
    "English (Canada) [en-CA]",
    "Spanish [es-ES]",
    "French [fr-FR]",
    "German [de-DE]",
    "Italian [it-IT]",
    "Portuguese (Brazil) [pt-BR]",
    "Japanese [ja-JP]",
    "Chinese (Simplified) [zh-CN]",
    "Custom / Other...",
    "None / Skip"
]

LANGUAGE_MAP = {
    "English (US) [en-US]": "en-US",
    "English (UK) [en-GB]": "en-GB",
    "English (Australia) [en-AU]": "en-AU",
    "English (Canada) [en-CA]": "en-CA",
    "Spanish [es-ES]": "es-ES",
    "French [fr-FR]": "fr-FR",
    "German [de-DE]": "de-DE",
    "Italian [it-IT]": "it-IT",
    "Portuguese (Brazil) [pt-BR]": "pt-BR",
    "Japanese [ja-JP]": "ja-JP",
    "Chinese (Simplified) [zh-CN]": "zh-CN",
    "None / Skip": ""
}

# ==========================================
# PART 1: STYLE SIMPLIFIER LOGIC
# ==========================================

def set_run_language(run, lang_code):
    rPr = run._element.get_or_add_rPr()
    for attr in ['w:val', 'w:eastAsia', 'w:bidi']:
        lang = OxmlElement('w:lang')
        lang.set(qn(attr), lang_code)
        rPr.append(lang)

def review_isolated_formatting(doc, root_window):
    """Reviews isolated formatted characters with prompt dialogs[cite: 1]."""
    for para in doc.paragraphs:
        current_pos = 0
        full_text = para.text
        for run in para.runs:
            clean_text = run.text.strip()
            run_len = len(run.text)

            if len(clean_text) == 1:
                active_formats = []
                if run.bold: active_formats.append("Bold")
                if run.italic: active_formats.append("Italic")
                if run.underline: active_formats.append("Underline")
                if run.font.strike: active_formats.append("Strikethrough")
                if run.font.name: active_formats.append(f"Font ({run.font.name})")
                if run.font.superscript: active_formats.append("Superscript")
                if run.font.subscript: active_formats.append("Subscript")
                if run.font.small_caps: active_formats.append("Small Caps")

                if active_formats:
                    start = max(0, current_pos - 30)
                    end = min(len(full_text), current_pos + 30)
                    before = full_text[start:current_pos]
                    after = full_text[current_pos + 1:end]
                    window = f"{before}[[{run.text}]]{after}"

                    msg = (
                        f"Isolated Formatting Detected:\n\n"
                        f"Context: ...{window}...\n"
                        f"Target: '{run.text}'\n"
                        f"Formatting: [{', '.join(active_formats)}]\n\n"
                        f"Do you want to KEEP this formatting? (Choose 'No' to strip it)"
                    )
                    keep = messagebox.askyesno("Isolated Formatting Review", msg, parent=root_window)
                    if not keep:
                        run.bold = run.italic = run.underline = run.font.strike = run.font.superscript = run.font.subscript = run.font.small_caps = False
                        run.font.name = None
            current_pos += run_len

def run_style_simplifier(input_file, output_file, style_map, config_options, lang_code="", root_window=None, log_callback=None):
    if not os.path.exists(input_file):
        if log_callback: log_callback(f"Error: File not found: {input_file}\n")
        return False

    doc = Document(input_file)
    if log_callback: log_callback("--- Running Style Simplifier Pass ---\n")

    strip_font = config_options.get("strip_font", False)
    strip_allcaps = config_options.get("strip_allcaps", False)
    strip_smallcaps = config_options.get("strip_smallcaps", False)
    strip_super_subscript = config_options.get("strip_super_subscript", False)
    strip_underline = config_options.get("strip_underline", False)
    remove_empty = config_options.get("remove_empty", False)
    strip_meta = config_options.get("strip_meta", False)
    review_isolated = config_options.get("review_isolated", False)

    for para in doc.paragraphs:
        if para.style.name in style_map:
            target_style = style_map[para.style.name]
            if target_style:
                try:
                    para.style = doc.styles[target_style]
                except KeyError:
                    try:
                        para.style = doc.styles[target_style.replace(" ", "")]
                    except KeyError:
                        pass

        pf = para.paragraph_format
        pf.line_spacing = pf.space_before = pf.space_after = pf.alignment = None
        pf.left_indent = pf.right_indent = pf.first_line_indent = None

        for run in para.runs:
            if run.font.hidden:
                run.text = ""
                continue

            run.style = None
            b, i_val, u, s = run.bold, run.italic, run.underline, run.font.strike
            rPr = run._element.get_or_add_rPr()
            tags_to_kill = [
                qn('w:sz'), qn('w:szCs'), qn('w:color'),
                qn('w:highlight'), qn('w:shd'), qn('w:cs')
            ]
            if strip_font: tags_to_kill.append(qn('w:rFonts'))
            if strip_allcaps: tags_to_kill.append(qn('w:caps'))
            if strip_smallcaps: tags_to_kill.append(qn('w:smallCaps'))
            if strip_super_subscript: tags_to_kill.append(qn('w:vertAlign'))
            if strip_underline:
                tags_to_kill.append(qn('w:u'))
                u = False

            for tag in tags_to_kill:
                element = rPr.find(tag)
                if element is not None:
                    rPr.remove(element)

            run.bold, run.italic, run.underline, run.font.strike = b, i_val, u, s
            if lang_code:
                set_run_language(run, lang_code)

    if remove_empty:
        if log_callback: log_callback("Removing empty paragraphs...\n")
        for para in list(doc.paragraphs):
            if not para.text.strip():
                p = para._element
                p.getparent().remove(p)

    if review_isolated:
        if log_callback: log_callback("Starting isolated character formatting review...\n")
        review_isolated_formatting(doc, root_window)

    if strip_meta:
        if log_callback: log_callback("Stripping document metadata...\n")
        core_props = doc.core_properties
        core_props.author = core_props.comments = core_props.keywords = ""
        core_props.last_modified_by = core_props.title = ""

    doc.save(output_file)
    if log_callback: log_callback(f"Cleaned document saved successfully to: {output_file}\n")
    return True


# ==========================================
# PART 2: DOCX2EPUB & MESSAGING CONVERTER LOGIC
# ==========================================

class CustomEpubHtml(epub.EpubHtml):
    def get_content(self):
        title_clean = self.title if self.title else "Chapter"
        html = (
            '<' + '?xml version="1.0" encoding="utf-8"?' + '>\n'
            '<' + '!DOCTYPE html>\n\n'
            '<' + 'html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" epub:prefix="z3998: http://www.daisy.org/z3998/2012/vocab/structure/#" lang="en" xml:lang="en">\n'
            '<' + 'head>\n'
            '<' + f'title>{title_clean}<' + '/title>\n'
            '  <' + 'link rel="stylesheet" type="text/css" href="styles/stylesheet.css"/>\n'
            '<' + 'link rel="stylesheet" type="text/css" href="styles/page_styles.css"/>\n'
            '<' + 'link rel="stylesheet" type="text/css" href="../fonts.css"/>\n'
            '<' + '/head>\n'
            '<' + 'body>\n'
            f'{self.content}\n'
            '<' + '/body>\n'
            '<' + '/html>'
        )
        return html.encode('utf-8')

def sanitize_style_name(name):
    clean = name.replace(' ', '_')
    clean = re.sub(r'[^a-zA-Z0-9_]', '', clean)
    return clean

def sanitize_filename_part(text):
    clean = re.sub(r'[^a-zA-Z0-9_]', '_', text)
    return re.sub(r'_+', '_', clean).strip('_')

def get_run_html(run):
    text = run.text
    if not text:
        return ""
    text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    classes = []
    inline_styles = []

    if run.bold: classes.append("bold")
    if run.italic: classes.append("italic")
    if run.underline: classes.append("underline")
    if run.font.strike: classes.append("strikethrough")
    if run.font.superscript: classes.append("superscript")
    if run.font.subscript: classes.append("subscript")
    if run.font.small_caps: classes.append("small-caps")
    if getattr(run.font, 'all_caps', False): classes.append("caps")
    if run.font.name:
        inline_styles.append(f"font-family: '{run.font.name}';")

    span_attrs = []
    if classes:
        span_attrs.append(f'class="{" ".join(classes)}"')
    if inline_styles:
        span_attrs.append(f'style="{" ".join(inline_styles)}"')

    if span_attrs:
        attr_str = " " + " ".join(span_attrs)
        text = f"<span{attr_str}>{text}</span>"

    return text

def prompt_sender_allocation(local_senders, root_window):
    """Opens a dialog window to allocate senders to message-me or message-them."""
    if not local_senders:
        return []

    # Simple custom dialog window for sender mapping
    dialog = ctk.CTk(fg_color="#2b2b2b") if not isinstance(root_window, ctk.CTk) else ctk.CTkToplevel(root_window)
    dialog.title("Allocate Message Senders")
    dialog.geometry("400x300")
    dialog.attributes("-topmost", True)
    dialog.grab_set()

    ctk.CTkLabel(dialog, text="Select who represents 'Message Me':", font=ctk.CTkFont(weight="bold")).pack(pady=10)
    ctk.CTkLabel(dialog, text="All other unselected senders will be 'Message Them'.", font=ctk.CTkFont(size=11), text_color="gray").pack(pady=2)

    selected_me_vars = {}
    frame = ctk.CTkScrollableFrame(dialog, height=140, width=350)
    frame.pack(padx=10, pady=10, fill="both", expand=True)

    for sender in local_senders:
        var = ctk.BooleanVar(value=False)
        chk = ctk.CTkCheckBox(frame, text=sender, variable=var)
        chk.pack(anchor="w", padx=10, pady=5)
        selected_me_vars[sender] = var

    me_names = []
    def on_confirm():
        nonlocal me_names
        me_names = [sender for sender, var in selected_me_vars.items() if var.get()]
        dialog.destroy()

    btn = ctk.CTkButton(dialog, text="Confirm Allocation", command=on_confirm)
    btn.pack(pady=10)

    dialog.wait_window()
    return me_names

def convert_message_batch_to_html(message_batch, root_window=None):
    days_of_week = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
    time_pattern = r'^\d{1,2}[:\.]\d{2}\s*(?:[aApP]\.?[mM]\.?)?$'

    local_senders = []
    for para in message_batch:
        raw_text = para.text.strip()
        if not raw_text or any(day == raw_text for day in days_of_week) or re.match(time_pattern, raw_text):
            continue
        if ":" in raw_text:
            name_part = raw_text.split(":", 1)[0].strip()
            if name_part and name_part not in local_senders:
                local_senders.append(name_part)

    # Prompt user to allocate senders interactively
    block_me_names = prompt_sender_allocation(local_senders, root_window)

    container_html = ['<div class="message-container">']
    current_time = ""
    last_sender = None
    new_timestamp_trigger = False

    for para in message_batch:
        raw_text = para.text.strip()
        if not raw_text or any(day == raw_text for day in days_of_week):
            continue

        if re.match(time_pattern, raw_text):
            current_time = raw_text
            new_timestamp_trigger = True
            continue

        if ":" in raw_text:
            name_part, message_content = raw_text.split(":", 1)
            name_part = name_part.strip()

            is_me = name_part in block_me_names
            row_class = "message-me-row" if is_me else "message-them-row"
            bubble_class = "message-me" if is_me else "message-them"
            name_class = "message-me-name" if is_me else "message-them-name"

            row_pieces = [f'<div class="{row_class}">']

            if name_part != last_sender or new_timestamp_trigger:
                ts = f" — {current_time}" if current_time else ""
                row_pieces.append(f'<div class="{name_class}">{name_part}{ts}</div>')
                new_timestamp_trigger = False

            content_html = ""
            colon_found = False
            for run in para.runs:
                run_text = run.text
                if not colon_found:
                    if ":" in run_text:
                        parts = run_text.split(":", 1)
                        colon_found = True
                        if parts[1]:
                            rem_text = parts[1]
                            formatted_rem = rem_text
                            if run.bold: formatted_rem = f'<span class="bold">{formatted_rem}</span>'
                            if run.italic: formatted_rem = f'<span class="italic">{formatted_rem}</span>'
                            if run.underline: formatted_rem = f'<span class="underline">{formatted_rem}</span>'
                            if run.font.strike: formatted_rem = f'<span class="strikethrough">{formatted_rem}</span>'
                            if run.font.superscript: formatted_rem = f'<span class="superscript">{formatted_rem}</span>'
                            if run.font.subscript: formatted_rem = f'<span class="subscript">{formatted_rem}</span>'
                            if run.font.small_caps: formatted_rem = f'<span class="small-caps">{formatted_rem}</span>'
                            if getattr(run.font, 'all_caps', False): formatted_rem = f'<span class="caps">{formatted_rem}</span>'
                            content_html += formatted_rem
                    continue
                else:
                    content_html += get_run_html(run)

            if not content_html.strip():
                content_html = message_content

            row_pieces.append(f'<div class="{bubble_class}">{content_html}</div>')
            row_pieces.append('</div>')
            container_html.append("".join(row_pieces))
            last_sender = name_part

    container_html.append('</div>')
    return "".join(container_html)

def generate_message_svg(container_html, epub_basename, chapter_title, section_number, stylesheet_path, output_dir="."):
    css_content = ""
    if stylesheet_path and os.path.exists(stylesheet_path):
        with open(stylesheet_path, 'r', encoding='utf-8') as sf:
            css_content = sf.read()

    svg_filename = os.path.join(output_dir, f"{sanitize_filename_part(epub_basename)}_{sanitize_filename_part(chapter_title)}_{section_number}.svg")

    svg_content = f'''<svg xmlns="http://www.w3.org/2000/svg" width="800" height="1000">
    <style>
        {css_content}
        body {{ background-color: #ffffff; font-family: sans-serif; }}
    </style>
    <foreignObject width="100%" height="100%">
        <div xmlns="http://www.w3.org/1999/xhtml">
            {container_html}
        </div>
    </foreignObject>
</svg>'''

    with open(svg_filename, 'w', encoding='utf-8') as svg_file:
        svg_file.write(svg_content)
    return svg_filename


# ==========================================
# CUSTOMTKINTER TWO-TAB INTERFACE
# ==========================================

class ConverterApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("Unified Docx Simplifier & EPUB/SVG Converter")
        self.geometry("950x850")
        self.minsize(850, 750)

        # Style Simplifier state variables
        self.simplifier_docx = ctk.StringVar()
        self.lang_var = ctk.StringVar(value="English (US) [en-US]")
        self.custom_lang_var = ctk.StringVar()
        self.style_map_widgets = {}
        self.style_checks = {}

        # EPUB & SVG Converter state variables
        self.epub_docx = ctk.StringVar()
        self.epub_output = ctk.StringVar()
        self.book_title = ctk.StringVar()
        self.book_author = ctk.StringVar()
        self.cover_path = ctk.StringVar()
        self.stylesheet_path = ctk.StringVar()
        self.pagestyles_path = ctk.StringVar()
        self.fontss_css_path = ctk.StringVar()
        self.svg_output_dir = ctk.StringVar(value=os.getcwd())

        self.gen_epub_var = ctk.BooleanVar(value=True)
        self.gen_svg_var = ctk.BooleanVar(value=True)

        self.font_css_mapping = {}
        self.extra_font_paths = []

        # Tabs Layout
        self.tabview = ctk.CTkTabview(self)
        self.tabview.pack(padx=20, pady=10, fill="both", expand=True)

        self.tab_simplifier = self.tabview.add("1. Style Simplifier")
        self.tab_epub = self.tabview.add("2. EPUB & SVG Generator")

        self.setup_simplifier_tab()
        self.setup_epub_tab()

        # Shared Console Log Box at bottom
        self.log_box = ctk.CTkTextbox(self, height=130)
        self.log_box.pack(padx=20, pady=10, fill="x")
        self.log_box.configure(state="disabled")

    def log_message(self, message):
        self.log_box.configure(state="normal")
        self.log_box.insert("end", message)
        self.log_box.see("end")
        self.log_box.configure(state="disabled")

    # ----------------------------------
    # TAB 1: STYLE SIMPLIFIER SETUP
    # ----------------------------------
    def setup_simplifier_tab(self):
        file_frame = ctk.CTkFrame(self.tab_simplifier)
        file_frame.pack(fill="x", padx=10, pady=10)

        lbl = ctk.CTkLabel(file_frame, text="Input DOCX:", width=100, anchor="w")
        lbl.grid(row=0, column=0, padx=5, pady=5, sticky="w")
        entry = ctk.CTkEntry(file_frame, textvariable=self.simplifier_docx, width=420)
        entry.grid(row=0, column=1, padx=5, pady=5)
        btn = ctk.CTkButton(file_frame, text="Browse", width=90, command=self.browse_simplifier_docx)
        btn.grid(row=0, column=2, padx=5, pady=5)

        # Language selection row
        lang_lbl = ctk.CTkLabel(file_frame, text="Language:", width=100, anchor="w")
        lang_lbl.grid(row=1, column=0, padx=5, pady=5, sticky="w")

        self.lang_dropdown = ctk.CTkOptionMenu(file_frame, values=LANGUAGE_OPTIONS, variable=self.lang_var, width=420, command=self.on_language_changed)
        self.lang_dropdown.grid(row=1, column=1, padx=5, pady=5, sticky="w")

        # Custom Language Frame
        self.custom_lang_frame = ctk.CTkFrame(file_frame, fg_color="transparent")
        self.custom_lang_frame.grid(row=2, column=1, padx=5, pady=2, sticky="w")
        self.custom_lang_lbl = ctk.CTkLabel(self.custom_lang_frame, text="Enter Lang Code (e.g., de-CH):", font=ctk.CTkFont(size=11))
        self.custom_lang_lbl.pack(side="left", padx=2)
        self.custom_lang_entry = ctk.CTkEntry(self.custom_lang_frame, textvariable=self.custom_lang_var, width=150)
        self.custom_lang_entry.pack(side="left", padx=2)
        self.custom_lang_frame.grid_remove()

        parse_btn = ctk.CTkButton(file_frame, text="Analyze Document Styles", fg_color="teal", hover_color="darkcyan", command=self.parse_document_styles)
        parse_btn.grid(row=3, column=1, padx=5, pady=10, sticky="w")

        # Scrollable container for dynamic style mapping & checkboxes
        self.simplifier_scroll = ctk.CTkScrollableFrame(self.tab_simplifier, height=240)
        self.simplifier_scroll.pack(fill="both", expand=True, padx=10, pady=5)

        run_simp_btn = ctk.CTkButton(self.tab_simplifier, text="Run Style Simplifier & Save Cleaned Docx", fg_color="green", hover_color="darkgreen", command=self.run_simplifier_action)
        run_simp_btn.pack(pady=10)

    def on_language_changed(self, choice):
        if choice == "Custom / Other...":
            self.custom_lang_frame.grid()
        else:
            self.custom_lang_frame.grid_remove()

    def browse_simplifier_docx(self):
        filename = filedialog.askopenfilename(title="Select Input DOCX", filetypes=[("Word Document", "*.docx")])
        if filename:
            self.simplifier_docx.set(filename)
            self.epub_docx.set(filename.replace(".docx", "_cleaned.docx"))

    def parse_document_styles(self):
        docx_path = self.simplifier_docx.get()
        if not docx_path or not os.path.exists(docx_path):
            messagebox.showerror("Error", "Please select a valid input DOCX file first.")
            return

        for widget in self.simplifier_scroll.winfo_children():
            widget.destroy()

        try:
            doc = Document(docx_path)
            used_styles = sorted(list({p.style.name for p in doc.paragraphs}))

            ctk.CTkLabel(self.simplifier_scroll, text="Style Mapping Configuration", font=ctk.CTkFont(weight="bold")).pack(anchor="w", pady=5)

            self.style_map_widgets = {}
            for name in used_styles:
                row_f = ctk.CTkFrame(self.simplifier_scroll, fg_color="transparent")
                row_f.pack(fill="x", pady=2)

                ctk.CTkLabel(row_f, text=f"Style: '{name}'", width=220, anchor="w").pack(side="left", padx=5)

                choice_var = ctk.StringVar(value="")
                dropdown = ctk.CTkOptionMenu(row_f, values=["Skip", "Heading 1", "Heading 2", "Normal"], variable=choice_var)
                dropdown.pack(side="left", padx=5)
                self.style_map_widgets[name] = choice_var

            ctk.CTkLabel(self.simplifier_scroll, text="\nFormatting Stripping Options", font=ctk.CTkFont(weight="bold")).pack(anchor="w", pady=5)

            self.style_checks = {
                "strip_allcaps": ctk.BooleanVar(value=False),
                "strip_smallcaps": ctk.BooleanVar(value=False),
                "strip_font": ctk.BooleanVar(value=False),
                "strip_super_subscript": ctk.BooleanVar(value=False),
                "strip_underline": ctk.BooleanVar(value=False),
                "remove_empty": ctk.BooleanVar(value=True),
                "strip_meta": ctk.BooleanVar(value=True),
                "review_isolated": ctk.BooleanVar(value=True)
            }

            labels_text = {
                "strip_allcaps": "Strip All-Caps formatting",
                "strip_smallcaps": "Strip Small-Caps formatting",
                "strip_font": "Strip custom font selections",
                "strip_super_subscript": "Strip Super/Subscript formatting",
                "strip_underline": "Strip Underline formatting",
                "remove_empty": "Remove empty paragraphs",
                "strip_meta": "Strip document metadata",
                "review_isolated": "Review isolated formatted single characters (e.g. typos)"
            }

            for key, bvar in self.style_checks.items():
                chk = ctk.CTkCheckBox(self.simplifier_scroll, text=labels_text[key], variable=bvar)
                chk.pack(anchor="w", padx=10, pady=2)

            self.log_message(f"Parsed styles for {os.path.basename(docx_path)} successfully.\n")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to parse document: {str(e)}")

    def run_simplifier_action(self):
        docx_path = self.simplifier_docx.get()
        if not docx_path or not os.path.exists(docx_path):
            messagebox.showerror("Error", "Valid input DOCX path required.")
            return

        output_file = filedialog.asksaveasfilename(defaultextension=".docx", title="Save Cleaned DOCX As", filetypes=[("Word Document", "*.docx")])
        if not output_file:
            return

        if os.path.exists(output_file):
            if not messagebox.askyesno("File Exists", f"'{os.path.basename(output_file)}' already exists. Do you want to overwrite it?"):
                return

        style_map = {}
        for style_name, var in self.style_map_widgets.items():
            val = var.get()
            if val != "Skip":
                style_map[style_name] = val

        config_options = {k: v.get() for k, v in self.style_checks.items()} if self.style_checks else {}
        selected_lang_label = self.lang_var.get()
        if selected_lang_label == "Custom / Other...":
            lang_code = self.custom_lang_var.get().strip()
        else:
            lang_code = LANGUAGE_MAP.get(selected_lang_label, "")

        threading.Thread(target=self._execute_simplifier, args=(docx_path, output_file, style_map, config_options, lang_code), daemon=True).start()

    def _execute_simplifier(self, src, out, smap, opts, lang_code):
        self.log_message("Starting Style Simplifier execution...\n")
        success = run_style_simplifier(src, out, smap, opts, lang_code=lang_code, root_window=self, log_callback=self.log_message)
        if success:
            self.epub_docx.set(out)
            messagebox.showinfo("Success", f"Cleaned document saved to:\n{out}")


    # ----------------------------------
    # TAB 2: EPUB & SVG GENERATOR SETUP
    # ----------------------------------
    def setup_epub_tab(self):
        scroll_tab2 = ctk.CTkScrollableFrame(self.tab_epub)
        scroll_tab2.pack(fill="both", expand=True, padx=5, pady=5)

        row_idx = 0
        self.add_file_row(scroll_tab2, "Cleaned DOCX:", self.epub_docx, lambda: self.browse_file(self.epub_docx, [("Word Document", "*.docx")]), row_idx); row_idx += 1
        self.add_file_row(scroll_tab2, "Output EPUB:", self.epub_output, self.browse_epub_output, row_idx); row_idx += 1

        mode_frame = ctk.CTkFrame(scroll_tab2, fg_color="transparent")
        mode_frame.grid(row=row_idx, column=0, columnspan=3, padx=5, pady=8, sticky="w")
        row_idx += 1

        ctk.CTkLabel(mode_frame, text="Output Targets:", font=ctk.CTkFont(weight="bold"), width=130, anchor="w").pack(side="left", padx=5)
        ctk.CTkCheckBox(mode_frame, text="Generate EPUB File", variable=self.gen_epub_var).pack(side="left", padx=15)
        ctk.CTkCheckBox(mode_frame, text="Generate Messaging SVGs", variable=self.gen_svg_var, command=self.toggle_svg_directory_visibility).pack(side="left", padx=15)

        lbl_title = ctk.CTkLabel(scroll_tab2, text="Book Title:", width=130, anchor="w")
        lbl_title.grid(row=row_idx, column=0, padx=5, pady=6, sticky="w")
        ent_title = ctk.CTkEntry(scroll_tab2, textvariable=self.book_title, width=400, placeholder_text="Optional (Defaults to filename)")
        ent_title.grid(row=row_idx, column=1, padx=5, pady=6, columnspan=2, sticky="w")
        row_idx += 1

        lbl_author = ctk.CTkLabel(scroll_tab2, text="Book Author:", width=130, anchor="w")
        lbl_author.grid(row=row_idx, column=0, padx=5, pady=6, sticky="w")
        ent_author = ctk.CTkEntry(scroll_tab2, textvariable=self.book_author, width=400, placeholder_text="Optional (Defaults to Unknown)")
        ent_author.grid(row=row_idx, column=1, padx=5, pady=6, columnspan=2, sticky="w")
        row_idx += 1

        self.add_file_row(scroll_tab2, "Cover Image:", self.cover_path, lambda: self.browse_file(self.cover_path, [("Images", "*.png *.jpg")]), row_idx); row_idx += 1
        self.add_file_row(scroll_tab2, "Stylesheet.css:", self.stylesheet_path, lambda: self.browse_file(self.stylesheet_path, [("CSS", "*.css")]), row_idx); row_idx += 1
        self.add_file_row(scroll_tab2, "Page Styles.css:", self.pagestyles_path, lambda: self.browse_file(self.pagestyles_path, [("CSS", "*.css")]), row_idx); row_idx += 1

        lbl = ctk.CTkLabel(scroll_tab2, text="Default Fonts.css:", width=130, anchor="w")
        lbl.grid(row=row_idx, column=0, padx=5, pady=6, sticky="w")
        entry = ctk.CTkEntry(scroll_tab2, textvariable=self.fontss_css_path, width=400)
        entry.grid(row=row_idx, column=1, padx=5, pady=6)
        btn = ctk.CTkButton(scroll_tab2, text="Browse", width=90, command=self.browse_font_css)
        btn.grid(row=row_idx, column=2, padx=5, pady=6)
        row_idx += 1

        self.svg_dir_label = ctk.CTkLabel(scroll_tab2, text="SVG Output Dir:", width=130, anchor="w")
        self.svg_dir_label.grid(row=row_idx, column=0, padx=5, pady=6, sticky="w")
        self.svg_dir_entry = ctk.CTkEntry(scroll_tab2, textvariable=self.svg_output_dir, width=400)
        self.svg_dir_entry.grid(row=row_idx, column=1, padx=5, pady=6)
        self.svg_dir_btn = ctk.CTkButton(scroll_tab2, text="Browse", width=90, command=self.browse_svg_directory)
        self.svg_dir_btn.grid(row=row_idx, column=2, padx=5, pady=6)
        row_idx += 1

        self.fonts_parse_frame = ctk.CTkFrame(scroll_tab2)
        self.fonts_parse_frame.grid(row=row_idx, column=0, columnspan=3, sticky="nsew", padx=5, pady=10)
        row_idx += 1

        ctk.CTkLabel(self.fonts_parse_frame, text="Required Fonts from fonts.css (Please map paths below):", font=ctk.CTkFont(weight="bold")).pack(anchor="w", padx=5, pady=5)
        self.fonts_mapping_inner = ctk.CTkFrame(self.fonts_parse_frame, fg_color="transparent")
        self.fonts_mapping_inner.pack(fill="x", padx=5, pady=5)

        self.extra_fonts_label = ctk.CTkLabel(scroll_tab2, text="Extra Release Fonts: 0 added", text_color="gray")
        self.extra_fonts_label.grid(row=row_idx, column=1, sticky="w", padx=5, pady=5)
        extra_btn = ctk.CTkButton(scroll_tab2, text="Add Extra Fonts", width=120, command=self.add_extra_fonts)
        extra_btn.grid(row=row_idx, column=2, padx=5, pady=5)
        row_idx += 1

        run_epub_btn = ctk.CTkButton(scroll_tab2, text="Run EPUB & SVG Workflow", fg_color="purple", hover_color="indigo", font=ctk.CTkFont(size=14, weight="bold"), command=self.run_epub_workflow)
        run_epub_btn.grid(row=row_idx, column=0, columnspan=3, pady=15)

    def add_file_row(self, parent, label_text, string_var, browse_command, row_idx):
        label = ctk.CTkLabel(parent, text=label_text, anchor="w", width=130)
        label.grid(row=row_idx, column=0, padx=5, pady=6, sticky="w")
        entry = ctk.CTkEntry(parent, textvariable=string_var, width=400)
        entry.grid(row=row_idx, column=1, padx=5, pady=6)
        btn = ctk.CTkButton(parent, text="Browse", width=90, command=browse_command)
        btn.grid(row=row_idx, column=2, padx=5, pady=6)

    def browse_file(self, string_var, filetypes):
        fn = filedialog.askopenfilename(title="Select File", filetypes=filetypes)
        if fn: string_var.set(fn)

    def browse_epub_output(self):
        fn = filedialog.asksaveasfilename(defaultextension=".epub", title="Save EPUB As", filetypes=[("EPUB", "*.epub")])
        if fn: self.epub_output.set(fn)

    def browse_svg_directory(self):
        dirname = filedialog.askdirectory(title="Select Directory for SVG Exports")
        if dirname:
            self.svg_output_dir.set(dirname)

    def toggle_svg_directory_visibility(self):
        if self.gen_svg_var.get():
            self.svg_dir_label.grid()
            self.svg_dir_entry.grid()
            self.svg_dir_btn.grid()
        else:
            self.svg_dir_label.grid_remove()
            self.svg_dir_entry.grid_remove()
            self.svg_dir_btn.grid_remove()

    def browse_font_css(self):
        fn = filedialog.askopenfilename(title="Select fonts.css", filetypes=[("CSS Files", "*.css")])
        if fn:
            self.fontss_css_path.set(fn)
            self.parse_fonts_css_file(fn)

    def parse_fonts_css_file(self, css_path):
        for widget in self.fonts_mapping_inner.winfo_children():
            widget.destroy()

        self.font_css_mapping = {}
        try:
            with open(css_path, 'r', encoding='utf-8') as f:
                content = f.read()

            urls = re.findall(r'url\s*\(\s*[\'"]?([^\'")\s]+)[\'"]?\s*\)', content)
            unique_fonts = sorted(list(set(urls)))

            if not unique_fonts:
                ctk.CTkLabel(self.fonts_mapping_inner, text="No font URLs detected in this CSS file.").pack(anchor="w")
                return

            for font_name in unique_fonts:
                row = ctk.CTkFrame(self.fonts_mapping_inner, fg_color="transparent")
                row.pack(fill="x", pady=2)

                ctk.CTkLabel(row, text=font_name, width=200, anchor="w").pack(side="left", padx=5)

                path_var = ctk.StringVar(value="")
                self.font_css_mapping[font_name] = path_var

                ent = ctk.CTkEntry(row, textvariable=path_var, width=250)
                ent.pack(side="left", padx=5)

                def make_browse(pvar):
                    return lambda: self.browse_font_file_path(pvar)

                b = ctk.CTkButton(row, text="Locate", width=70, command=make_browse(path_var))
                b.pack(side="left", padx=5)

            self.log_message(f"Parsed {len(unique_fonts)} font reference(s) from fonts.css\n")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to parse fonts.css: {str(e)}")

    def browse_font_file_path(self, string_var):
        fn = filedialog.askopenfilename(title="Select Font File", filetypes=[("Font Files", "*.ttf *.otf *.woff *.woff2")])
        if fn: string_var.set(fn)

    def add_extra_fonts(self):
        filenames = filedialog.askopenfilenames(title="Select Extra Release Fonts", filetypes=[("Font Files", "*.ttf *.otf *.woff *.woff2")])
        if filenames:
            self.extra_font_paths.extend(filenames)
            self.extra_fonts_label.configure(text=f"Extra Release Fonts: {len(self.extra_font_paths)} added", text_color="green")

    def run_epub_workflow(self):
        if not self.gen_epub_var.get() and not self.gen_svg_var.get():
            messagebox.showerror("Error", "Please select at least one output option (Generate EPUB or Generate SVGs).")
            return

        docx_path = self.epub_docx.get()
        if not docx_path or not os.path.exists(docx_path):
            messagebox.showerror("Error", "Please provide a valid cleaned DOCX file.")
            return

        output_epub = self.epub_output.get()
        if not output_epub:
            output_epub = docx_path.replace('.docx', '.epub')

        if self.gen_epub_var.get() and os.path.exists(output_epub):
            if not messagebox.askyesno("File Exists", f"EPUB file '{os.path.basename(output_epub)}' already exists. Overwrite?"):
                return

        threading.Thread(target=self._execute_epub_conversion, args=(docx_path, output_epub), daemon=True).start()

    def _execute_epub_conversion(self, docx_path, output_path):
        self.log_message("Starting conversion workflow...\n")
        try:
            doc = Document(docx_path)
            epub_basename = os.path.splitext(os.path.basename(output_path))[0]
            svg_out_dir = self.svg_output_dir.get() or "."
            if self.gen_svg_var.get() and not os.path.exists(svg_out_dir):
                os.makedirs(svg_out_dir, exist_ok=True)

            book = epub.EpubBook()
            book.set_identifier(str(uuid.uuid4()))

            final_title = self.book_title.get().strip() or epub_basename
            final_author = self.book_author.get().strip() or "Unknown Author"

            book.set_title(final_title)
            book.set_language('en')
            book.add_author(final_author)

            cover = self.cover_path.get()
            if cover and os.path.exists(cover):
                ext = os.path.splitext(cover)[1].lower()
                with open(cover, 'rb') as f:
                    book.set_cover("cover" + ext, f.read())

            stylesheet = self.stylesheet_path.get()
            css_bytes = b""
            if stylesheet and os.path.exists(stylesheet):
                with open(stylesheet, 'rb') as f:
                    css_bytes = f.read()

            additional_css = (
                b"/* Additional document styles generated by Antithesis Press epub generator */\n"
                b".bold { font-weight: bold; }\n"
                b".italic { font-style: italic; }\n"
                b".underline { text-decoration: underline; }\n"
                b".strikethrough { text-decoration: line-through; }\n"
                b".superscript { vertical-align: super; font-size: smaller; }\n"
                b".subscript { vertical-align: sub; font-size: smaller; }\n"
                b".small-caps { font-variant: small-caps; }\n"
                b".caps { text-transform: uppercase; }\n\n"
            )
            final_stylesheet_content = additional_css + css_bytes

            book.add_item(epub.EpubItem(uid="stylesheet", file_name="content/styles/stylesheet.css", media_type="text/css", content=final_stylesheet_content))

            pagestyles = self.pagestyles_path.get()
            if pagestyles and os.path.exists(pagestyles):
                with open(pagestyles, 'rb') as f:
                    book.add_item(epub.EpubItem(uid="page_styles", file_name="content/styles/page_styles.css", media_type="text/css", content=f.read()))

            font_css = self.fontss_css_path.get()
            if font_css and os.path.exists(font_css):
                with open(font_css, 'rb') as f:
                    book.add_item(epub.EpubItem(uid="fonts_css", file_name="fonts.css", media_type="text/css", content=f.read()))

            all_fonts_to_include = list(self.extra_font_paths)
            for fname_key, path_var in self.font_css_mapping.items():
                real_path = path_var.get()
                if real_path and os.path.exists(real_path):
                    all_fonts_to_include.append(real_path)

            for fpath in all_fonts_to_include:
                if os.path.exists(fpath):
                    fname = os.path.basename(fpath)
                    ext = fname.split('.')[-1].lower()
                    media_type = f"font/{ext}" if ext in ['woff', 'woff2', 'ttf', 'otf'] else "application/octet-stream"
                    with open(fpath, 'rb') as f:
                        book.add_item(epub.EpubItem(uid=fname, file_name=fname, media_type=media_type, content=f.read()))

            chapters = []
            current_chapter_content = []
            current_chapter_title = "Preface"
            message_section_counters = {}

            paragraphs = doc.paragraphs
            i = 0
            while i < len(paragraphs):
                para = paragraphs[i]
                style_name = para.style.name

                if style_name.lower() == 'messages':
                    message_batch = []
                    while i < len(paragraphs) and paragraphs[i].style.name.lower() == 'messages':
                        message_batch.append(paragraphs[i])
                        i += 1

                    # Pass self (root window) for interactive dialog allocation
                    container_html = convert_message_batch_to_html(message_batch, root_window=self)
                    current_chapter_content.append(container_html)

                    if current_chapter_title not in message_section_counters:
                        message_section_counters[current_chapter_title] = 0
                    message_section_counters[current_chapter_title] += 1
                    sec_num = message_section_counters[current_chapter_title]

                    if self.gen_svg_var.get():
                        svg_filename = generate_message_svg(container_html, epub_basename, current_chapter_title, sec_num, stylesheet, svg_out_dir)
                        self.log_message(f"Generated SVG: {os.path.basename(svg_filename)}\n")
                    continue

                text = para.text.strip()
                if not text:
                    i += 1
                    continue

                tag = 'div'
                css_class = ""
                if style_name == 'Normal':
                    tag = 'p'
                elif style_name == 'Heading 1':
                    tag = 'h1'
                else:
                    tag = 'div'
                    css_class = f' class="{sanitize_style_name(style_name)}"'

                inner_html = "".join([get_run_html(run) for run in para.runs])

                if tag == 'h1':
                    if current_chapter_content:
                        chap_file = f"content/chapter_{len(chapters)}.xhtml"
                        chapters.append(CustomEpubHtml(title=current_chapter_title, file_name=chap_file, lang='en', content="".join(current_chapter_content)))

                    current_chapter_title = para.text.strip()
                    current_chapter_content = []
                    current_chapter_content.append(f"<h1>{inner_html}</h1>")
                else:
                    current_chapter_content.append(f"<{tag}{css_class}>{inner_html}</{tag}>")

                i += 1

            if current_chapter_content:
                chap_file = f"content/chapter_{len(chapters)}.xhtml"
                chapters.append(CustomEpubHtml(title=current_chapter_title, file_name=chap_file, lang='en', content="".join(current_chapter_content)))

            if self.gen_epub_var.get():
                def make_page(title, filename, body_text):
                    page = CustomEpubHtml(title=title, file_name=f"content/{filename}")
                    page.content = f'<div style="text-align:center; margin-top:30%;"><h1>{title}</h1><p>{body_text}</p></div>'
                    return page

                cover_page = make_page('Cover Page', 'cover.xhtml', 'cover page')
                title_page = make_page('Title Page', 'title.xhtml', 'title page')
                author_page = make_page('About the Author', 'about_author.xhtml', 'about the author')
                press_page = make_page('About the Press', 'about_press.xhtml', 'about the press')
                copyright_page = make_page('Copyright Information', 'copyright.xhtml', 'insert copyright information')

                book.add_item(cover_page)
                book.add_item(title_page)
                for chap in chapters: book.add_item(chap)
                book.add_item(author_page)
                book.add_item(press_page)
                book.add_item(copyright_page)

                book.toc = chapters
                book.spine = ['nav', cover_page, title_page] + chapters + [author_page, press_page, copyright_page]
                book.add_item(epub.EpubNcx())
                book.add_item(epub.EpubNav())

                self.log_message(f"Writing EPUB file to {output_path}...\n")
                epub.write_epub(output_path, book, {})

            self.log_message("Workflow completed successfully!\n")
            messagebox.showinfo("Success", "Selected output tasks finished successfully!")
        except Exception as e:
            self.log_message(f"[ERROR]: {str(e)}\n")
            messagebox.showerror("Error", f"An error occurred: {str(e)}")


if __name__ == "__main__":
    app = ConverterApp()
    app.mainloop()
